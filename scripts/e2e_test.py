"""端到端测试：样例发票 → 起服务 → 上传（含 Tracker 基准）→ 融合导出校验。

用法：python scripts/e2e_test.py
"""
from __future__ import annotations

import http.cookiejar
import json
import shutil
import subprocess
import sys
import time
import urllib.error
import urllib.request
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.utils.units import pixels_to_EMU

COOKIE_JAR = http.cookiejar.CookieJar()
OPENER = urllib.request.build_opener(urllib.request.HTTPCookieProcessor(COOKIE_JAR))

ROOT = Path(__file__).resolve().parent.parent
SAMPLES = Path(__file__).parent / "samples"
FIXTURES = Path(__file__).parent / "fixtures"
BLANK_TRACKER = FIXTURES / "tracker_blank.xlsx"     # 只有 Tracker list（用户真实基准的副本）
FULL_TRACKER = FIXTURES / "tracker_full.xlsx"       # 已有 Expenses report + Bills 的账本
# 独立数据目录 + 随机空闲端口：绝不触碰真实会话数据
E2E_DATA = Path(__file__).parent / "e2e_data"


def free_port() -> int:
    import socket
    s = socket.socket()
    s.bind(("127.0.0.1", 0))
    port = s.getsockname()[1]
    s.close()
    return port


PORT = free_port()
BASE = f"http://127.0.0.1:{PORT}"

EMU_PER_PX = 9525
HALF_COL_PX = 9 * 64
passed, failed = 0, 0


def check(name: str, cond: bool, detail: str = ""):
    global passed, failed
    if cond:
        passed += 1
        print(f"  ✓ {name}")
    else:
        failed += 1
        print(f"  ✗ {name}  {detail}")


def req(method: str, path: str, data=None, raw=None, headers=None):
    url = BASE + path
    body = raw
    hdr = dict(headers or {})
    if data is not None:
        body = json.dumps(data).encode()
        hdr["Content-Type"] = "application/json"
    r = urllib.request.Request(url, data=body, headers=hdr, method=method)
    try:
        with OPENER.open(r, timeout=120) as resp:
            return resp.read()
    except urllib.error.HTTPError as e:
        return json.dumps({"__status__": e.code, **json.loads(e.read() or b"{}")}).encode()


def login(username: str, password: str) -> tuple[int, dict]:
    body = json.dumps({"username": username, "password": password}).encode()
    r = urllib.request.Request(BASE + "/api/login", data=body,
                               headers={"Content-Type": "application/json"}, method="POST")
    try:
        with OPENER.open(r, timeout=60) as resp:
            return resp.status, json.loads(resp.read().decode())
    except urllib.error.HTTPError as e:
        return e.code, json.loads(e.read() or b"{}")


def multipart_upload(files: list[Path], field: str = "files") -> dict:
    boundary = "----zcode-e2e-boundary"
    body = b""
    for p in files:
        body += (
            f"--{boundary}\r\n"
            f'Content-Disposition: form-data; name="{field}"; filename="{p.name}"\r\n'
            f"Content-Type: application/octet-stream\r\n\r\n"
        ).encode() + p.read_bytes() + b"\r\n"
    body += f"--{boundary}--\r\n".encode()
    resp = req(
        "POST", "/api/upload" if field == "files" else "/api/tracker", raw=body,
        headers={"Content-Type": f"multipart/form-data; boundary={boundary}"},
    )
    return json.loads(resp)


def make_full_tracker():
    """伪造一本已有内容的账本：Expenses report 已有 2 行，Bills 已有 1#、2#。"""
    shutil.copy(BLANK_TRACKER, FULL_TRACKER)
    wb = load_workbook(FULL_TRACKER)
    ws = wb.create_sheet("Expenses report")
    ws.merge_cells("A1:F1")
    ws["A1"] = "ZENDO - Travelling Expenses report"
    ws["A1"].font = ws["A1"].font.copy(size=18)
    for j, h in enumerate(["Date:", "Item:", "Location:", "Amt (RMB):", "HK$", "Bill No."], 1):
        c = ws.cell(row=2, column=j, value=h)
        c.font = c.font.copy(bold=True)
    import datetime as dt
    ws["A3"] = dt.datetime(2026, 9, 10)
    ws["A3"].number_format = "d-mmm"
    ws["B3"], ws["C3"], ws["D3"], ws["F3"] = "滴滴打车", "旧记录一", 50, "1#"
    ws["A4"] = dt.datetime(2026, 9, 10)
    ws["A4"].number_format = "d-mmm"
    ws["B4"], ws["C4"], ws["D4"], ws["F4"] = "高铁", "旧记录二", 80, "2#"
    ws2 = wb.create_sheet("Bills")
    ws2["E2"] = "1#"
    ws2["N2"] = "2#"
    wb.save(FULL_TRACKER)


def image_x_ranges(ws) -> list[tuple[int, int, int]]:
    """返回每张图的 (起始x_px, 结束x_px, 锚点行)。"""
    out = []
    for img in ws._images:
        a = img.anchor
        x0 = a._from.col * 64 + a._from.colOff / EMU_PER_PX
        w = a.ext.cx / EMU_PER_PX
        out.append((x0, x0 + w, a._from.row))
    return out


def main() -> int:
    subprocess.run([sys.executable, str(Path(__file__).parent / "make_samples.py")], check=True)
    if not BLANK_TRACKER.exists():
        shutil.copy(r"C:\Users\Admin\Desktop\Tracker list.xlsx", BLANK_TRACKER)
    make_full_tracker()

    server = subprocess.Popen(
        [sys.executable, "-m", "uvicorn", "app.main:app", "--host", "127.0.0.1", "--port", str(PORT)],
        cwd=str(ROOT / "backend"),
        env={
            **__import__("os").environ,
            "INVOICE_DATA_DIR": str(E2E_DATA),
        },
        stdout=open(ROOT / "scripts" / "server.log", "w", encoding="utf-8"),
        stderr=subprocess.STDOUT,
    )
    try:
        for _ in range(60):
            try:
                urllib.request.urlopen(BASE + "/api/me", timeout=2)
                break
            except urllib.error.URLError:
                time.sleep(0.5)
        else:
            print("服务启动失败")
            return 1

        files = sorted(SAMPLES.iterdir())

        # ================================================= 0. 登录与安全
        print("== 0. 登录与安全 ==")
        body = req("GET", "/api/invoices")
        check("未登录访问 API 返回 401", json.loads(body).get("__status__") == 401, body.decode()[:80])
        hint = (E2E_DATA / "初始密码.txt").read_text(encoding="utf-8")
        import re as _re
        m = _re.search(r"用户名:\s*(\S+)\s*密码:\s*(\S+)", hint)
        check("自动生成初始账号", m is not None, hint[:60])
        USER, PW = (m.group(1), m.group(2)) if m else ("admin", "")
        code, out = login(USER, "wrong-password-x")
        check("错误密码返回 401", code == 401 and "用户名或密码错误" in out.get("detail", ""),
              f"{code} {out}")
        code, out = login(USER, PW)
        check("正确密码登录成功", code == 200 and out.get("ok") is True, f"{code} {out}")
        me = json.loads(req("GET", "/api/me").decode())
        check("会话生效 /api/me", me.get("logged_in") is True and me.get("username") == USER, str(me))

        # ================================================= A. 干净 Tracker 基准（只有 Tracker list 表）→ 融合导出
        print("== A. 干净 Tracker 基准（只有 Tracker list 表）→ 融合导出 ==")
        req("POST", "/api/session/clear")
        req("DELETE", "/api/tracker")
        up = multipart_upload([BLANK_TRACKER], field="file")
        check("基准已加载", up["loaded"] is True and "Tracker list" in up["sheets"], str(up))
        added = multipart_upload(files)
        check("5 张发票解析", len(added) == 5, str(len(added)))
        by = {a["filename"]: a for a in added}
        check("赣州站-长沙站 → 赣州站到长沙站",
              by["rail-ganzhou-changsha.pdf"]["location"] == "赣州站到长沙站",
              by["rail-ganzhou-changsha.pdf"]["location"])
        check("高铁 Location=A站到B站", by["rail-0912.pdf"]["location"] == "泉州南站到长汀南站",
              by["rail-0912.pdf"]["location"])
        check("火车 Location=A站到B站（补站名）", by["train-0919.pdf"]["location"] == "赣州站到兴国站",
              by["train-0919.pdf"]["location"])
        check("滴滴 Location 留空待人工", by["didi-0911.pdf"]["location"] == ""
              and by["didi-0922.png"]["location"] == "",
              f'{by["didi-0911.pdf"]["location"]!r}/{by["didi-0922.png"]["location"]!r}')
        req("POST", "/api/invoices/manual", data={
            "date": "2026-09-12", "item": "高铁",
            "location": "泉州南至长汀南站", "amount": 0, "note": "（工厂已付）",
        })

        out = ROOT / "scripts" / "export_test.xlsx"
        out.write_bytes(req("GET", "/api/export"))
        wb = load_workbook(out)
        check("三个 sheet", wb.sheetnames == ["Tracker list", "Expenses report", "Bills"],
              str(wb.sheetnames))
        check("Tracker list 原样保留", wb["Tracker list"]["A4"].value == "Vendor",
              str(wb["Tracker list"]["A4"].value))
        ws = wb["Expenses report"]
        headers = [ws.cell(row=2, column=j).value for j in range(1, 7)]
        check("表头一致", headers == ["Date:", "Item:", "Location:", "Amt (RMB):", "HK$", "Bill No."],
              str(headers))
        rows = [r for r in ws.iter_rows(min_row=3, values_only=True) if any(v is not None for v in r)]
        check("路程 6 行", len(rows) == 6, str(len(rows)))
        check("编号 1#-5#", [r[5] for r in rows if r[5] and str(r[5]).endswith("#")]
              == ["1#", "2#", "3#", "4#", "5#"], str([r[5] for r in rows]))
        ws2 = wb["Bills"]
        labels = [c.value for row in ws2.iter_rows() for c in row
                  if isinstance(c.value, str) and c.value.endswith("#")]
        check("Bills 编号 1#-5#", sorted(labels) == [f"{i}#" for i in range(1, 6)], str(labels))
        check("5 张图嵌入", len(ws2._images) == 5, str(len(ws2._images)))
        xr = image_x_ranges(ws2)
        lefts = [x for x in xr if x[0] < HALF_COL_PX]
        rights = [x for x in xr if x[0] >= HALF_COL_PX]
        check("左图不越中线", all(x[1] <= HALF_COL_PX for x in lefts),
              str([(round(a), round(b)) for a, b, _ in xr]))
        check("右图不越中线", all(x[0] >= HALF_COL_PX for x in rights), "")
        check("左右不重叠（逐对）", not any(
            not (x1[1] <= x0[0] or x0[1] <= x1[0]) and not (x1[2] + 20 < x0[2] or x0[2] + 20 < x1[2])
            for i, x0 in enumerate(xr) for x1 in xr[i + 1:]), str(sorted(xr)))
        first = ws2._images[0].anchor
        check("第一张锚 A 列、标签下、半栏居中",
              first._from.col == 0 and first._from.row == 2 and 0 < first._from.colOff < pixels_to_EMU(HALF_COL_PX),
              f"col={first._from.col} row={first._from.row} colOff={first._from.colOff}")
        wpx = first.ext.cx / EMU_PER_PX
        check("图宽 ≤570px", 100 < wpx <= 570, str(wpx))

        # ================================================= B. 已有账本追加
        print("== B. 已有 Expenses/Bills 的账本 → 追加编号衔接 ==")
        req("POST", "/api/session/clear")
        multipart_upload([FULL_TRACKER], field="file")
        multipart_upload(files)
        out.write_bytes(req("GET", "/api/export"))
        wb = load_workbook(out)
        check("三个 sheet", wb.sheetnames == ["Tracker list", "Expenses report", "Bills"],
              str(wb.sheetnames))
        ws = wb["Expenses report"]
        vals = []
        for row in ws.iter_rows(min_row=3, values_only=True):
            if any(v is not None for v in row):
                vals.append(row)
        check("旧 2 行 + 新 5 行 = 7", len(vals) == 7, str(len(vals)))
        check("旧行原样在前面", vals[0][5] == "1#" and vals[1][5] == "2#"
              and vals[0][2] == "旧记录一", str(vals[0]))
        new_nos = [r[5] for r in vals[2:] if r[5]]
        check("新行编号 3#-7# 衔接", new_nos == ["3#", "4#", "5#", "6#", "7#"], str(new_nos))
        ws2 = wb["Bills"]
        labels = sorted(
            (c.value for row in ws2.iter_rows() for c in row
             if isinstance(c.value, str) and c.value.endswith("#")),
            key=lambda s: int(s[:-1]))
        check("编号 1#-7#", labels == [f"{i}#" for i in range(1, 8)], str(labels))
        lab3 = [c for row in ws2.iter_rows(min_row=20, max_row=30) for c in row if c.value == "3#"]
        check("新标签 3# 在 E25 块", lab3 and lab3[0].coordinate == "E25",
              lab3[0].coordinate if lab3 else "无")
        xr = image_x_ranges(ws2)
        check("追加后 5 张新图", len(xr) == 5, str(len(xr)))
        rows_expect = {3: 25, 4: 25, 5: 48, 6: 48, 7: 71}
        ok_rows = all(a[2] == rows_expect[i + 3] for i, a in enumerate(sorted(xr, key=lambda t: t[2])))
        check("新图块行位置 25/25/48/48/71", ok_rows, str(sorted(xr, key=lambda t: t[2])))
        lefts = [x for x in xr if x[0] < HALF_COL_PX]
        rights = [x for x in xr if x[0] >= HALF_COL_PX]
        check("追加后仍无重叠", all(x[1] <= HALF_COL_PX for x in lefts)
              and all(x[0] >= HALF_COL_PX for x in rights), str(sorted(xr)))

        # ================================================= C. 无基准导出
        print("== C. 未上传基准 → 仅两个 sheet ==")
        req("DELETE", "/api/tracker")
        out.write_bytes(req("GET", "/api/export"))
        wb = load_workbook(out)
        check("两个 sheet", wb.sheetnames == ["Expenses report", "Bills"], str(wb.sheetnames))

        # ================================================= C2. 登录锁定（放最后，会锁住账号）
        print("== C2. 连续失败锁定 ==")
        codes = []
        for i in range(5):
            c2, _o = login(USER, "bad-pass")
            codes.append(c2)
        check("连续失败均 401", all(c == 401 for c in codes), str(codes))
        code, out = login(USER, PW)   # 密码正确也拒绝
        check("达到上限后锁定", code == 401 and "锁定" in out.get("detail", ""), f"{code} {out}")

        print("== D. 前端静态页 ==")
        html = req("GET", "/")
        check("首页可访问", b"<div id=\"app\">" in html)

        print(f"\n结果: {passed} 通过, {failed} 失败")
        return 1 if failed else 0
    finally:
        server.terminate()
        try:
            server.wait(timeout=10)
        except subprocess.TimeoutExpired:
            server.kill()


if __name__ == "__main__":
    sys.exit(main())
